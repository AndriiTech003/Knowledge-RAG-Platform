from __future__ import annotations

import io
from pathlib import Path

import pymupdf
from docx import Document

HERE = Path(__file__).parent
HEADER = "Northwind Labs · Field Operations Manual"


def hard_pdf() -> bytes:
    doc = pymupdf.open()
    pages = [
        "<h1>Field Operations Manual</h1><p>This manual describes how field engineers prepare site visits. "
        "Every visit needs a signed work order before the engineer leaves the office.</p>"
        "<h2>Preparation</h2><p>Engineers check the equipment list, confirm the customer contact and book a vehicle "
        "through the fleet portal at least two working days in advance.</p>",
        "<h2>Equipment allowance</h2><p>The table lists the standard kit per visit type.</p>"
        "<table><tr><th>Visit type</th><th>Kit</th><th>Max weight</th></tr>"
        "<tr><td>Install</td><td>Kit A</td><td>18 kg</td></tr>"
        "<tr><td>Repair</td><td>Kit B</td><td>12 kg</td></tr>"
        "<tr><td>Audit</td><td>Kit C</td><td>4 kg</td></tr></table>"
        "<p>Heavier kits must be split across two engineers. The safety officer for heavy kits is Mira Kovac.</p>",
        "<h2>Reporting</h2><p>After the visit, the engineer files report FR-209 within 24 hours. Reports that are "
        "late are escalated to the regional lead, who reviews them every Friday.</p>",
    ]
    for index, html in enumerate(pages, start=1):
        page = doc.new_page(width=595, height=842)
        page.insert_text((56, 34), HEADER, fontsize=8, fontname="helv")
        page.insert_text(
            (56, 814), f"Confidential - Page {index} of {len(pages)}", fontsize=8, fontname="helv"
        )
        page.insert_htmlbox(
            pymupdf.Rect(56, 60, 539, 790),
            html,
            css="h1{font-size:20pt} h2{font-size:15pt} p,td,th{font-size:10.5pt} "
            "table{border-collapse:collapse} td,th{border:1px solid #888; padding:2pt}",
        )
    doc.set_metadata({"title": "Field Operations Manual"})
    data = doc.tobytes()
    doc.close()
    return bytes(data)


def scanned_pdf() -> bytes:
    doc = pymupdf.open()
    for _ in range(2):
        page = doc.new_page(width=595, height=842)
        pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 200), False)
        pix.clear_with(200)
        page.insert_image(pymupdf.Rect(50, 50, 400, 400), pixmap=pix)
    data = doc.tobytes()
    doc.close()
    return bytes(data)


def sample_docx() -> bytes:
    document = Document()
    document.core_properties.title = "Office Moves Guide"
    document.add_heading("Office Moves Guide", level=1)
    document.add_paragraph("Moves are coordinated by the workplace team two weeks ahead.")
    document.add_heading("Packing", level=2)
    document.add_paragraph("Each employee receives three crates.", style="List Bullet")
    document.add_paragraph("Monitors are packed by the movers.", style="List Bullet")
    document.add_heading("Costs", level=2)
    table = document.add_table(rows=3, cols=2)
    for r, row in enumerate([("Item", "Cost"), ("Crates", "$40"), ("Movers", "$1,200")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


SAMPLE_MD = """# Coffee Machine Guide

The kitchen coffee machine is a Lungo 3000.

## Cleaning

- Descale every **Monday**.
- Empty the grounds drawer daily.

## Supplies

| Item | Supplier |
|---|---|
| Beans | Roastery Nine |
| Milk | Dairy Co |
"""

SAMPLE_HTML = """<!doctype html><html><head><title>Parking Rules | Intranet</title></head><body>
<nav><a href="/">Home</a><a href="/policies">Policies</a></nav>
<header>Intranet header banner</header>
<main><article><h1>Parking Rules</h1><p>The garage opens at 6:30 and closes at 22:00 on weekdays.</p>
<h2>Electric vehicles</h2><p>Charging bays are limited to four hours per car per day.</p>
<ul><li>Bay A: 11 kW</li><li>Bay B: 22 kW</li></ul></article></main>
<footer>Copyright footer text that should be removed</footer></body></html>
"""

SAMPLE_TXT = "Visitor Wi-Fi\n\nThe visitor network is NW-Visit.\n\nPasswords rotate weekly.\n"


def main() -> None:
    (HERE / "hard.pdf").write_bytes(hard_pdf())
    (HERE / "scanned.pdf").write_bytes(scanned_pdf())
    (HERE / "sample.docx").write_bytes(sample_docx())
    (HERE / "sample.md").write_text(SAMPLE_MD)
    (HERE / "sample.html").write_text(SAMPLE_HTML)
    (HERE / "sample.txt").write_text(SAMPLE_TXT)


if __name__ == "__main__":
    main()
