from __future__ import annotations

import argparse
import html
import shutil
import sys
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument
from docx.shared import Pt
from markdown_it import MarkdownIt

PAGE_BREAK = "<<<PAGE>>>"
PAGE_WIDTH = 595
PAGE_HEIGHT = 842
MARGIN = 56

PDF_CSS = """
body { font-family: sans-serif; font-size: 10.5pt; line-height: 1.35; }
h1 { font-size: 20pt; margin: 0 0 8pt 0; }
h2 { font-size: 15pt; margin: 10pt 0 6pt 0; }
h3 { font-size: 12.5pt; margin: 8pt 0 4pt 0; }
p { margin: 0 0 6pt 0; }
table { border-collapse: collapse; margin: 4pt 0 8pt 0; }
th, td { border: 1px solid #888; padding: 2pt 5pt; font-size: 9.5pt; }
li { margin: 0 0 2pt 0; }
code, pre { font-family: monospace; font-size: 9pt; }
"""

HTML_TEMPLATE = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>{title} | Northwind Labs Intranet</title>
<link rel="canonical" href="https://intranet.northwind.example/{slug}">
</head>
<body>
<header class="site-header"><a href="/">Northwind Labs Intranet</a>
<nav><ul><li><a href="/handbook">Handbook</a></li><li><a href="/policies">Policies</a></li>
<li><a href="/engineering">Engineering</a></li><li><a href="/search">Search</a></li></ul></nav></header>
<div class="cookie-banner">We use cookies to improve the intranet. <button>Accept</button></div>
<main><article>
{body}
</article></main>
<aside class="related"><h4>Related pages</h4><ul><li><a href="/handbook/welcome">Welcome</a></li></ul></aside>
<footer class="site-footer">Copyright 2026 Northwind Labs. All rights reserved. Intranet v4.2</footer>
</body>
</html>
"""


@dataclass
class SourceDoc:
    path: Path
    collection: str
    name: str
    title: str
    fmt: str
    meta: dict[str, str] = field(default_factory=dict)
    body: str = ""


def parse_source(path: Path, root: Path) -> SourceDoc:
    raw = path.read_text(encoding="utf-8")
    meta: dict[str, str] = {}
    body = raw
    if raw.startswith("---\n"):
        end = raw.index("\n---", 4)
        for line in raw[4:end].splitlines():
            if ":" in line:
                key, value = line.split(":", 1)
                meta[key.strip()] = value.strip()
        body = raw[end + 4 :].lstrip("\n")
    rel = path.relative_to(root)
    collection = rel.parts[0]
    fmt = meta.get("format", "md")
    title = meta.get("title", path.stem.replace("-", " ").title())
    return SourceDoc(
        path=path, collection=collection, name=path.stem, title=title, fmt=fmt, meta=meta, body=body
    )


def render_pdf(doc: SourceDoc, out: Path) -> int:
    md = MarkdownIt("commonmark").enable("table")
    pages = [p.strip() for p in doc.body.split(PAGE_BREAK)]
    pdf = pymupdf.open()
    total = len(pages)
    for index, page_md in enumerate(pages, start=1):
        page = pdf.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
        header = f"Northwind Labs · {doc.title}"
        page.insert_text((MARGIN, 34), header, fontsize=8, fontname="helv", color=(0.4, 0.4, 0.4))
        footer = f"Confidential - Internal use only · Page {index} of {total}"
        page.insert_text(
            (MARGIN, PAGE_HEIGHT - 28), footer, fontsize=8, fontname="helv", color=(0.4, 0.4, 0.4)
        )
        rect = pymupdf.Rect(MARGIN, 52, PAGE_WIDTH - MARGIN, PAGE_HEIGHT - 48)
        spare, _scale = page.insert_htmlbox(rect, md.render(page_md), css=PDF_CSS, scale_low=1)
        if spare < 0:
            raise ValueError(f"{doc.path}: page {index} overflows, split it with {PAGE_BREAK}")
    pdf.set_metadata({"title": doc.title, "author": doc.meta.get("author", "Northwind Labs")})
    pdf.save(out)
    pdf.close()
    return total


def _add_runs(paragraph: object, children: list[object]) -> None:
    from docx.text.paragraph import Paragraph

    assert isinstance(paragraph, Paragraph)
    bold = False
    italic = False
    for child in children:
        kind = getattr(child, "type", "")
        if kind == "strong_open":
            bold = True
        elif kind == "strong_close":
            bold = False
        elif kind == "em_open":
            italic = True
        elif kind == "em_close":
            italic = False
        elif kind in {"text", "code_inline"}:
            run = paragraph.add_run(str(getattr(child, "content", "")))
            run.bold = bold
            run.italic = italic
        elif kind in {"softbreak", "hardbreak"}:
            paragraph.add_run(" ")


def render_docx(doc: SourceDoc, out: Path) -> None:
    md = MarkdownIt("commonmark").enable("table")
    tokens = md.parse(doc.body.replace(PAGE_BREAK, ""))
    document = DocxDocument()
    style = document.styles["Normal"]
    font = getattr(style, "font", None)
    if font is not None:
        font.size = Pt(11)
    document.core_properties.title = doc.title
    document.core_properties.author = doc.meta.get("author", "Northwind Labs")
    i = 0
    list_style: list[str] = []
    while i < len(tokens):
        tok = tokens[i]
        if tok.type == "heading_open":
            level = int(tok.tag[1])
            inline = tokens[i + 1]
            document.add_heading(inline.content, level=min(level, 3) if level > 1 else 1)
            i += 3
            continue
        if tok.type == "bullet_list_open":
            list_style.append("List Bullet")
        elif tok.type == "ordered_list_open":
            list_style.append("List Number")
        elif tok.type in {"bullet_list_close", "ordered_list_close"}:
            list_style.pop()
        elif tok.type == "paragraph_open":
            inline = tokens[i + 1]
            style_name = list_style[-1] if list_style else None
            paragraph = document.add_paragraph(style=style_name)
            _add_runs(paragraph, list(inline.children or []))
            i += 3
            continue
        elif tok.type in {"fence", "code_block"}:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(tok.content.rstrip("\n"))
            run.font.name = "Courier New"
        elif tok.type == "table_open":
            rows: list[list[str]] = []
            j = i + 1
            current: list[str] = []
            while tokens[j].type != "table_close":
                if tokens[j].type == "tr_open":
                    current = []
                elif tokens[j].type == "inline":
                    current.append(tokens[j].content)
                elif tokens[j].type == "tr_close":
                    rows.append(current)
                j += 1
            width = max(len(r) for r in rows)
            table = document.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for r_index, row in enumerate(rows):
                for c_index, cell in enumerate(row):
                    table.cell(r_index, c_index).text = cell
            i = j + 1
            continue
        i += 1
    document.save(str(out))


def render_html(doc: SourceDoc, out: Path) -> None:
    md = MarkdownIt("commonmark").enable("table")
    body = md.render(doc.body.replace(PAGE_BREAK, ""))
    slug = f"{doc.collection}/{doc.name}"
    out.write_text(HTML_TEMPLATE.format(title=html.escape(doc.title), slug=slug, body=body), encoding="utf-8")


def render_md(doc: SourceDoc, out: Path) -> None:
    out.write_text(doc.body.replace(PAGE_BREAK, "").strip() + "\n", encoding="utf-8")


EXTENSIONS = {"pdf": ".pdf", "docx": ".docx", "html": ".html", "md": ".md", "txt": ".txt"}


def build(src_root: Path, out_root: Path) -> list[tuple[str, str, int | None]]:
    if out_root.exists():
        shutil.rmtree(out_root)
    out_root.mkdir(parents=True)
    built: list[tuple[str, str, int | None]] = []
    for path in sorted(src_root.rglob("*.md")):
        doc = parse_source(path, src_root)
        target_dir = out_root / doc.collection
        target_dir.mkdir(parents=True, exist_ok=True)
        target = target_dir / f"{doc.name}{EXTENSIONS[doc.fmt]}"
        pages: int | None = None
        if doc.fmt == "pdf":
            pages = render_pdf(doc, target)
        elif doc.fmt == "docx":
            render_docx(doc, target)
        elif doc.fmt == "html":
            render_html(doc, target)
        elif doc.fmt == "txt":
            target.write_text(doc.body.replace(PAGE_BREAK, ""), encoding="utf-8")
        else:
            render_md(doc, target)
        built.append((f"{doc.collection}/{target.name}", doc.fmt, pages))
    return built


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="kb-corpus")
    base = Path(__file__).resolve().parents[4] / "eval-data"
    parser.add_argument("--src", type=Path, default=base / "corpus-src")
    parser.add_argument("--out", type=Path, default=base / "corpus")
    args = parser.parse_args(argv)
    built = build(args.src, args.out)
    for name, fmt, pages in built:
        sys.stdout.write(f"{name}\t{fmt}\t{pages if pages is not None else '-'}\n")
    sys.stdout.write(f"{len(built)} documents\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
